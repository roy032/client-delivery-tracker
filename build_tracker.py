"""Builds Client_Delivery_Tracker.xlsx: a request log, a formula-driven dashboard and a weekly report sheet.

Run: python build_tracker.py   (needs openpyxl)
All data is sample data for demonstration.
"""
import csv
from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F3864"
hdr_font = Font(bold=True, color="FFFFFF", name="Calibri")
hdr_fill = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

TYPES = ["Software issue", "Data request", "Documentation", "Training", "New feature", "Report"]
STATUS = ["Open", "In progress", "Waiting on client", "Resolved", "Closed"]
PRIORITY = ["High", "Medium", "Low"]
SLA_DAYS = {"High": 2, "Medium": 5, "Low": 10}  # target days to resolve

rows = list(csv.DictReader(open("sample_requests.csv", encoding="utf-8")))

wb = Workbook()

# ---------- Settings ----------
st = wb.active
st.title = "Settings"
st["A1"], st["B1"] = "Priority", "SLA target (days)"
for i, (p, d) in enumerate(SLA_DAYS.items(), start=2):
    st.cell(i, 1, p); st.cell(i, 2, d)
st["D1"] = "Report date (change to today)"
st["E1"] = date(2026, 10, 12)
st["E1"].number_format = "DD-MMM-YYYY"
for c in ("A1", "B1", "D1"):
    st[c].font = hdr_font; st[c].fill = hdr_fill
st.column_dimensions["A"].width = 14; st.column_dimensions["B"].width = 18
st.column_dimensions["D"].width = 30; st.column_dimensions["E"].width = 16

# ---------- Requests ----------
ws = wb.create_sheet("Requests")
heads = ["ID", "Client", "Request type", "Priority", "Opened", "Status", "Closed on",
         "Days open", "SLA target (days)", "SLA status", "Notes", "Follow-up flag"]
ws.append(heads)
for i, r in enumerate(rows, start=2):
    opened = date.fromisoformat(r["opened"])
    closed = date.fromisoformat(r["closed"]) if r["closed"] else None
    ws.append([r["id"], r["client"], r["type"], r["priority"], opened, r["status"], closed,
               f'=IF(G{i}="",Settings!$E$1-E{i},G{i}-E{i})',
               f'=VLOOKUP(D{i},Settings!$A$2:$B$4,2,FALSE)',
               f'=IF(H{i}<=I{i},IF(G{i}="","On track","Met"),IF(G{i}="","Overdue","Missed"))',
               r["notes"], f'=IF(J{i}="Overdue",A{i}&" ("&B{i}&")","")'])
n = len(rows) + 1
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.border = box
    row[4].number_format = "DD-MMM-YYYY"; row[6].number_format = "DD-MMM-YYYY"
tab = Table(displayName="RequestsTable", ref=f"A1:L{n}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tab)
for col, w in zip("ABCDEFGHIJKL", [8, 22, 16, 10, 13, 18, 13, 10, 16, 12, 44, 26]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
for rng, lst in (("C2:C500", TYPES), ("D2:D500", PRIORITY), ("F2:F500", STATUS)):
    dv = DataValidation(type="list", formula1='"' + ",".join(lst) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)
red = PatternFill("solid", bgColor="F8CBAD"); green = PatternFill("solid", bgColor="C6E0B4")
ws.conditional_formatting.add(f"J2:J{n}", CellIsRule(operator="equal", formula=['"Overdue"'], fill=red))
ws.conditional_formatting.add(f"J2:J{n}", CellIsRule(operator="equal", formula=['"Missed"'], fill=red))
ws.conditional_formatting.add(f"J2:J{n}", CellIsRule(operator="equal", formula=['"Met"'], fill=green))

# ---------- Dashboard ----------
db = wb.create_sheet("Dashboard", 0)
db["A1"] = "Client Delivery Dashboard"; db["A1"].font = Font(bold=True, size=16, color=NAVY)
db["A2"] = '="As of "&TEXT(Settings!E1,"DD MMM YYYY")'
kp = [("Total requests", f"=COUNTA(Requests!A2:A{n})"),
      ("Still open", f'=COUNTA(Requests!A2:A{n})-COUNTIF(Requests!F2:F{n},"Closed")-COUNTIF(Requests!F2:F{n},"Resolved")'),
      ("Overdue", f'=COUNTIF(Requests!J2:J{n},"Overdue")'),
      ("SLA met (closed)", f'=IFERROR(COUNTIF(Requests!J2:J{n},"Met")/(COUNTIF(Requests!J2:J{n},"Met")+COUNTIF(Requests!J2:J{n},"Missed")),0)'),
      ("Avg days to resolve", f'=IFERROR(AVERAGEIF(Requests!G2:G{n},"<>",Requests!H2:H{n}),0)')]
for i, (k, f) in enumerate(kp, start=4):
    db.cell(i, 1, k).font = Font(bold=True); db.cell(i, 2, f)
db["B7"].number_format = "0%"; db["B8"].number_format = "0.0"
db["A11"], db["B11"], db["C11"] = "Status", "Count", ""
for i, s in enumerate(STATUS, start=12):
    db.cell(i, 1, s); db.cell(i, 2, f"=COUNTIF(Requests!F2:F{n},A{i})")
db["D11"], db["E11"] = "Request type", "Count"
for i, t in enumerate(TYPES, start=12):
    db.cell(i, 4, t); db.cell(i, 5, f"=COUNTIF(Requests!C2:C{n},D{i})")
db["G11"], db["H11"] = "Client", "Open requests"
clients = sorted({r["client"] for r in rows})
for i, c in enumerate(clients, start=12):
    db.cell(i, 7, c)
    db.cell(i, 8, f'=COUNTIFS(Requests!B2:B{n},G{i},Requests!F2:F{n},"<>Closed",Requests!F2:F{n},"<>Resolved")')
for c in ("A11", "B11", "D11", "E11", "G11", "H11"):
    db[c].font = hdr_font; db[c].fill = hdr_fill
for col, w in zip("ABCDEFGH", [22, 12, 3, 18, 10, 3, 24, 14]):
    db.column_dimensions[col].width = w
ch = BarChart(); ch.type = "col"; ch.title = "Requests by status"; ch.legend = None
ch.add_data(Reference(db, min_col=2, min_row=11, max_row=11 + len(STATUS)), titles_from_data=True)
ch.set_categories(Reference(db, min_col=1, min_row=12, max_row=11 + len(STATUS)))
ch.height, ch.width = 7, 13
db.add_chart(ch, "A20")

# ---------- Weekly report ----------
wr = wb.create_sheet("Weekly report")
wr["A1"] = "Weekly report (copy into Word / email)"; wr["A1"].font = Font(bold=True, size=14, color=NAVY)
wr["A3"] = '="Summary: "&Dashboard!B4&" requests logged in total, "&Dashboard!B5&" still open, "&Dashboard!B6&" overdue."'
wr["A4"] = '="SLA met on "&TEXT(Dashboard!B7,"0%")&" of closed requests; average resolution time "&TEXT(Dashboard!B8,"0.0")&" days."'
wr["A6"] = "Overdue items needing follow-up:"; wr["A6"].font = Font(bold=True)
wr["A7"] = '=IF(COUNTIF(Requests!L2:L%d,"?*")=0,"none",_xlfn.TEXTJOIN(", ",TRUE,Requests!L2:L%d))' % (n, n)
wr.column_dimensions["A"].width = 120
wb.save("Client_Delivery_Tracker.xlsx")
print("saved", len(rows), "requests")
